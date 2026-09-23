#!/usr/bin/env python3
"""Exercise conversion/XML regressions; optionally pass a contract Markdown path."""
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET

from ooxml_order import ORDERS

SCRIPTS = Path(__file__).resolve().parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def run(script, *args, expected=0, diagnostic='Full ordered text mismatch'):
    result = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)],
                            capture_output=True, text=True)
    assert result.returncode == expected, result.stdout + result.stderr
    if expected:
        assert diagnostic in result.stdout, result.stdout
    return result.stdout


def check(path, source, expected=0):
    run('check-a4-docx.py', path, source, expected=expected)
    run('check-docx-typography.py', path, '--fonts',
        'Pretendard|Pretendard Medium|Pretendard SemiBold', '--sizes', '10.5,12,16',
        '--min-size', '10.5', '--no-bold', '--exact-sizes', '--source', source,
        expected=expected)


def convert(source, workspace, inject_spacing=False):
    base, final = workspace / 'base.docx', workspace / 'final.docx'
    run('md-to-a4-docx.py', source, '-o', base)
    run('check-a4-docx.py', base, source)
    if inject_spacing:
        doc = Document(base)
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:val'), '20')
        doc.paragraphs[0].runs[0]._r.get_or_add_rPr().append(spacing)
        doc.save(base)
    run('apply-docx-typography.py', base, '-o', final,
        '--regular-face', 'Pretendard', '--medium-face', 'Pretendard Medium',
        '--semibold-face', 'Pretendard SemiBold', '--monochrome')
    check(final, source)
    with zipfile.ZipFile(final) as archive:
        counts = {'rPr': 0, 'pPr': 0}
        for name in archive.namelist():
            if not (name.startswith('word/') and name.endswith('.xml')):
                continue
            root = ET.fromstring(archive.read(name))
            for node in root.iter():
                if node.tag not in (qn('w:rPr'), qn('w:pPr')):
                    continue
                ranks = [ORDERS[node.tag][n.tag] for n in node if n.tag in ORDERS[node.tag]]
                assert ranks == sorted(ranks), (name, ET.tostring(node))
                counts[ET.QName(node).localname] += 1
        print(f'PASS: all Word XML property orders {counts}')
        styles = ET.fromstring(archive.read('word/styles.xml'))
        order = ['name', 'aliases', 'basedOn', 'next', 'link', 'autoRedefine',
                 'hidden', 'uiPriority', 'semiHidden', 'unhideWhenUsed', 'qFormat',
                 'locked', 'personal', 'personalCompose', 'personalReply', 'rsid',
                 'pPr', 'rPr', 'tblPr', 'trPr', 'tcPr', 'tblStylePr']
        for style in styles.findall('w:style', NS):
            ranks = [order.index(ET.QName(node).localname) for node in style]
            assert ranks == sorted(ranks), ET.tostring(style)
            if style.get(qn('w:type')) == 'character':
                assert style.find('w:pPr', NS) is None
        if inject_spacing:
            body = ET.fromstring(archive.read('word/document.xml'))
            spacing = body.findall('.//w:rPr/w:spacing', NS)
            assert len(spacing) == 1
            assert dict(spacing[0].attrib) == {qn('w:val'): '20'}
            siblings = [ET.QName(n).localname for n in spacing[0].getparent()]
            assert siblings.index('color') < siblings.index('spacing') < siblings.index('sz')
            assert siblings.index('rFonts') < siblings.index('b') < siblings.index('color')
    print('PASS: both checkers, CT_Style order, character style properties')
    return final


def reject_bad_order(final, workspace):
    for prop, first, last in (('rPr', 'rFonts', 'sz'), ('pPr', 'spacing', 'rPr')):
        bad_path = workspace / f'bad-{prop}-order.docx'
        with zipfile.ZipFile(final) as source, zipfile.ZipFile(bad_path, 'w') as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == 'word/document.xml':
                    root = ET.fromstring(data)
                    parent = next(n for n in root.iter(qn(f'w:{prop}'))
                                  if n.find(qn(f'w:{first}')) is not None
                                  and n.find(qn(f'w:{last}')) is not None)
                    node = parent.find(qn(f'w:{first}'))
                    parent.remove(node)
                    parent.append(node)
                    data = ET.tostring(root)
                target.writestr(info, data)
        output = run('check-docx-typography.py', bad_path, '--fonts',
                     'Pretendard|Pretendard Medium|Pretendard SemiBold',
                     '--sizes', '10.5,12,16', '--min-size', '10.5', '--no-bold',
                     expected=1, diagnostic='Property order violation')
        assert 'word/document.xml' in output and prop in output
        print(f'PASS: reversed {prop} order rejected by typography checker (exit 1)')


def main():
    with tempfile.TemporaryDirectory(prefix='a4-regressions-', dir=SCRIPTS.parent.parent) as tmp:
        workspace = Path(tmp)
        source = workspace / 'synthetic.md'
        source.write_text('# Real title\n\n## Real section\n\n```md\n# Fake title\n'
                          '| Fake | Table |\n| --- | --- |\n| a | b |\n```\n\n'
                          '| Real | Table |\n| --- | --- |\n| a | b |\n\n'
                          '상호: __________ / 대표자: __________\n\n'
                          '_기울임_ __굵게__ *기울임* **굵게** [링크](https://example.com)<br>다음 줄\n')
        final = convert(source, workspace, inject_spacing=True)
        reject_bad_order(final, workspace)
        doc = Document(final)
        assert len(doc.tables) == 1
        assert sum(p.style.name.startswith('Heading') for p in doc.paragraphs) == 2
        assert '# Fake title' in '\n'.join(p.text for p in doc.paragraphs)
        print('PASS: fenced examples remain literal, real counts 2 headings/1 table; run spacing val=20 unchanged')
        if len(sys.argv) > 1:
            source = Path(sys.argv[1])
            final = convert(source, workspace)
            doc = Document(final)
            expected = '상호/성명: __________ / 대표자: __________ / 주소: __________ / 사업자등록번호: __________ / 연락처: __________ / 세금계산서 이메일: __________'
            assert doc.tables[0].rows[1].cells[1].text == expected
            # Count physical w:t nodes once.
            body_text = ''.join(n.text or '' for n in doc.element.body.iter(qn('w:t')))
            assert body_text.count('_') == source.read_text().count('_') == 118
            print('PASS: contract first information cell exact; underscores=118; tables=', len(doc.tables))
        bad = Document(final)
        run_node = bad.paragraphs[0].runs[0]
        run_node.text = 'X' + run_node.text[1:]
        bad_path = workspace / 'bad.docx'
        bad.save(bad_path)
        check(bad_path, source, expected=1)
        print('PASS: single-character mutation rejected by both checkers (exit 1)')
    assert not workspace.exists()
    print('PASS: temporary workspace removed')


if __name__ == '__main__':
    main()
