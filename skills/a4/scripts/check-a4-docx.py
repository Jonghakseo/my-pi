#!/usr/bin/env python3
"""DOCX sanity checks for /a4 output."""
from __future__ import annotations

import re
import html
import sys
import zipfile
from pathlib import Path

from markdown_structure import outside_code_fences

from docx import Document


def markdown_table_count(markdown: str) -> int:
    lines = outside_code_fences(markdown).splitlines()
    count = 0
    for i in range(len(lines) - 1):
        if lines[i].strip().startswith("|") and re.match(r"^\|\s*:?-{3,}:?", lines[i + 1].strip()):
            count += 1
    return count


def markdown_heading_count(markdown: str) -> dict[str, int]:
    counts = {f"h{i}": 0 for i in range(1, 7)}
    for m in re.finditer(r"^(#{1,6})\s+\S.*$", outside_code_fences(markdown), flags=re.M):
        counts[f"h{len(m.group(1))}"] += 1
    return counts


def docx_heading_count(doc: Document) -> dict[str, int]:
    counts = {f"h{i}": 0 for i in range(1, 7)}
    for p in doc.paragraphs:
        name = p.style.name if p.style else ""
        m = re.fullmatch(r"Heading ([1-6])", name)
        if m:
            counts[f"h{m.group(1)}"] += 1
    return counts


def markdown_visible_text(markdown: str) -> str:
    """Extract source text in order, stripping only supported Markdown syntax.

    Code spans/blocks are literal. Ordered-list numbers and contract blanks
    remain text. This reader does not call the DOCX renderer.
    """
    inline = re.compile(
        r"`([^`]+)`|\*\*([^*]+)\*\*|"
        r"(?<![\w_])__([^_\s](?:[^_]*?[^_\s])?)__(?![\w_])|"
        r"(?<!\*)\*([^*]+)\*(?!\*)|"
        r"(?<![\w_])_([^_\s](?:[^_]*?[^_\s])?)_(?![\w_])|"
        r"\[([^\]]+)\]\([^)]+\)"
    )

    def visible(text: str) -> str:
        text = text.replace("<br />", "\n").replace("<br/>", "\n").replace("<br>", "\n")
        return html.unescape(inline.sub(
            lambda match: next(group for group in match.groups() if group is not None), text
        ))

    lines = markdown.splitlines()
    result = []
    in_code = False
    in_table = False
    for index, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            in_table = False
            continue
        if in_code:
            result.append(line)
            continue
        table_row = stripped.startswith("|") and stripped.endswith("|")
        if not table_row:
            in_table = False
        if table_row and index + 1 < len(lines):
            cells = lines[index + 1].strip().strip("|").split("|")
            if all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells):
                in_table = True
        if in_table and table_row:
            cells = stripped.strip("|").split("|")
            if not all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells):
                result.extend(visible(cell.strip()) for cell in cells)
            continue
        if re.fullmatch(r"-{3,}|_{3,}|\*{3,}", stripped):
            continue
        line = re.sub(r"^#{1,6}\s+", "", line)
        line = re.sub(r"^\s*[-*+]\s+", "", line)
        line = re.sub(r"^\s*>+\s*", "", line)
        result.append(visible(line))
    return "\n".join(result)


def docx_ordered_text(doc: Document) -> str:
    """Read body paragraphs and table cells in XML order, excluding footers."""
    from docx.oxml.ns import qn

    # XML order traverses each physical cell once, including nested tables.
    parts = []
    for node in doc.element.body.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag in (qn("w:tab"), qn("w:br"), qn("w:cr"), qn("w:p")):
            parts.append("\n")
    return "".join(parts)


def check_full_ordered_text(markdown: str, doc: Document) -> str | None:
    """Return a first-difference diagnostic, or None for matching visible text."""
    expected = re.sub(r"\s+", "", markdown_visible_text(markdown))
    actual = re.sub(r"\s+", "", docx_ordered_text(doc))
    if expected == actual:
        return None
    index = next((i for i, (a, b) in enumerate(zip(expected, actual)) if a != b),
                 min(len(expected), len(actual)))
    start, end = max(0, index - 30), index + 40
    return (f"Full ordered text mismatch at character {index + 1} "
            f"(whitespace ignored; source={len(expected)}, DOCX={len(actual)} chars):\n"
            f"  source[{start}:{end}]: {expected[start:end]!r}\n"
            f"  DOCX  [{start}:{end}]: {actual[start:end]!r}")


def main() -> int:
    if len(sys.argv) not in (2, 3):
        print("Usage: check-a4-docx.py output.docx [source.md]", file=sys.stderr)
        return 2

    docx_path = Path(sys.argv[1]).expanduser().resolve()
    md_path = Path(sys.argv[2]).expanduser().resolve() if len(sys.argv) == 3 else None
    errors: list[str] = []

    doc = Document(docx_path)
    visible_text = docx_ordered_text(doc)

    if re.search(r"<\s*br\s*/?\s*>|&lt;\s*br\s*/?\s*&gt;", visible_text, re.I):
        errors.append("Visible DOCX text contains literal <br>. It must be an actual line break, not text.")

    non_one = []
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() and p.paragraph_format.line_spacing not in (None, 1.0):
            non_one.append(idx)
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    if p.text.strip() and p.paragraph_format.line_spacing not in (None, 1.0):
                        non_one.append(f"t{ti}r{ri}c{ci}")
    if non_one:
        errors.append(f"Found non-1.0 line spacing overrides: {non_one[:10]}")

    if md_path:
        md = md_path.read_text(encoding="utf-8")
        text_error = check_full_ordered_text(md, doc)
        if text_error:
            errors.append(text_error)
        md_headings = markdown_heading_count(md)
        doc_headings = docx_heading_count(doc)
        if md_headings != doc_headings:
            errors.append(f"Heading count mismatch: markdown={md_headings}, docx={doc_headings}")
        md_tables = markdown_table_count(md)
        if md_tables != len(doc.tables):
            errors.append(f"Table count mismatch: markdown={md_tables}, docx={len(doc.tables)}")

    with zipfile.ZipFile(docx_path) as z:
        bad = z.testzip()
        if bad:
            errors.append(f"DOCX zip integrity failed at {bad}")
        footer_xml = "\n".join(z.read(n).decode("utf-8", errors="ignore") for n in z.namelist() if n.startswith("word/footer") and n.endswith(".xml"))
        document_xml = z.read("word/document.xml")
        styles_xml = z.read("word/styles.xml") if "word/styles.xml" in z.namelist() else b""
        if "PAGE" not in footer_xml:
            errors.append("Footer PAGE field is missing.")
        if b"w:numPr" in document_xml:
            errors.append("Unexpected Word automatic numbering (w:numPr) found.")
        if any(token in document_xml + styles_xml for token in [b"w:keepNext", b"w:keepLines"]):
            errors.append("Unexpected keepNext/keepLines found.")

    if errors:
        print("A4 DOCX check: FAIL")
        for err in errors:
            print(f"- {err}")
        return 1

    print("A4 DOCX check: OK")
    if md_path:
        print("Full ordered text: OK")
    print(f"Tables: {len(doc.tables)}")
    print("Footer PAGE field: OK")
    print("Visible <br> text: none")
    print("Base line spacing: 1.0")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
